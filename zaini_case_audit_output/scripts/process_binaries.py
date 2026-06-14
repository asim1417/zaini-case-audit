#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
process_binaries.py
معالجة الملفات الثنائية المعقّدة محلياً (دون MCP):
  - فكّ ضغط zip/rar/7z من staging/raw/<id>.<ext> إلى staging/extracted/<id>/
  - قراءة الملفات الداخلية (xlsx/pdf/docx/txt/csv) + OCR للصور (tesseract ara+eng)
  - كتابة النص المجمّع لكل أرشيف/صورة إلى staging/text/<id>.txt
  - إنتاج جرد للملفات الداخلية في staging/extracted/_inner_inventory.json
الاستخدام: python3 process_binaries.py <staging_dir>
"""
import sys, os, re, json, subprocess, zipfile
from pathlib import Path

STAGING = Path(sys.argv[1])
RAW = STAGING / "raw"
EXC = STAGING / "extracted"
TEXT = STAGING / "text"
EXC.mkdir(parents=True, exist_ok=True)
TEXT.mkdir(parents=True, exist_ok=True)

IMG_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp"}


def ocr_image(path):
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(path)
        return pytesseract.image_to_string(img, lang="ara+eng")
    except Exception as e:  # noqa
        return f"[OCR error: {e}]"


def ocr_pdf(path, max_pages=40):
    """OCR لملف PDF صورة (بلا طبقة نص) عبر تحويل الصفحات إلى صور ثم tesseract."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
        pages = convert_from_path(str(path), dpi=200)
        out = []
        for i, pg in enumerate(pages[:max_pages], 1):
            out.append(f"[صفحة {i} - OCR]")
            out.append(pytesseract.image_to_string(pg, lang="ara+eng"))
        return "\n".join(out)
    except Exception as e:  # noqa
        return f"[PDF OCR error: {e}]"


def read_inner(path):
    ext = path.suffix.lower()
    try:
        if ext in (".txt", ".csv"):
            return path.read_text(encoding="utf-8", errors="replace")
        if ext == ".docx":
            import docx
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        if ext in (".xlsx", ".xlsm", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"# ورقة: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" | ".join(cells))
            return "\n".join(out)
        if ext == ".pdf":
            import pypdf
            r = pypdf.PdfReader(str(path))
            txt = "\n".join((pg.extract_text() or "") for pg in r.pages)
            if len(txt.strip()) < 20:  # PDF صورة → OCR للصفحات
                return ocr_pdf(path)
            return txt
        if ext in IMG_EXT:
            return ocr_image(path)
        return ""
    except Exception as e:  # noqa
        return f"[read error {path.name}: {e}]"


def extract_archive(raw_path, dest):
    dest.mkdir(parents=True, exist_ok=True)
    ext = raw_path.suffix.lower()
    try:
        if ext == ".zip":
            try:
                with zipfile.ZipFile(raw_path) as z:
                    z.extractall(dest)
                return True, "zip"
            except Exception:
                subprocess.run(["7z", "x", "-y", f"-o{dest}", str(raw_path)],
                               capture_output=True, check=True)
                return True, "zip/7z"
        if ext == ".rar":
            r = subprocess.run(["unrar", "x", "-o+", str(raw_path), str(dest) + "/"],
                               capture_output=True)
            if r.returncode != 0:
                subprocess.run(["7z", "x", "-y", f"-o{dest}", str(raw_path)],
                               capture_output=True, check=True)
            return True, "rar"
        if ext in (".7z",):
            subprocess.run(["7z", "x", "-y", f"-o{dest}", str(raw_path)],
                           capture_output=True, check=True)
            return True, "7z"
    except Exception as e:  # noqa
        return False, f"extract error: {e}"
    return False, "unknown ext"


inner_inventory = {}
processed = 0
if not RAW.exists():
    print("لا يوجد مجلد raw - لم يجرِ تنزيل ملفات ثنائية بعد.")
    sys.exit(0)

for raw_path in sorted(RAW.glob("*")):
    if raw_path.is_dir():
        continue
    fid = raw_path.stem
    ext = raw_path.suffix.lower()
    out_txt = TEXT / f"{fid}.txt"

    if ext in IMG_EXT:
        text = ocr_image(raw_path)
        out_txt.write_text(f"[OCR لصورة {raw_path.name} عبر tesseract ara+eng]\n\n{text}",
                           encoding="utf-8")
        inner_inventory[fid] = {"type": "image", "chars": len(text)}
        processed += 1
        continue

    if ext in (".zip", ".rar", ".7z"):
        dest = EXC / fid
        ok, how = extract_archive(raw_path, dest)
        if not ok:
            out_txt.write_text(f"[تعذّر فكّ الأرشيف {raw_path.name}: {how}]", encoding="utf-8")
            inner_inventory[fid] = {"type": "archive", "extracted": False, "note": how}
            continue
        inner_files = [p for p in dest.rglob("*") if p.is_file()]
        parts = [f"[أرشيف {raw_path.name} - عدد الملفات الداخلية: {len(inner_files)} - طريقة: {how}]", ""]
        inv = []
        for p in sorted(inner_files):
            rel = str(p.relative_to(dest))
            t = read_inner(p)
            inv.append({"name": rel, "ext": p.suffix.lower(),
                        "size": p.stat().st_size, "chars": len(t or "")})
            parts.append(f"\n===== ملف داخلي: {rel} =====")
            parts.append(t or "[لا نص]")
        out_txt.write_text("\n".join(parts), encoding="utf-8")
        inner_inventory[fid] = {"type": "archive", "extracted": True,
                                "inner_count": len(inner_files), "inner": inv}
        processed += 1
        continue

    # ملفات أخرى (pdf/docx/...) نزّلت خاماً
    text = read_inner(raw_path)
    out_txt.write_text(text or "[لا نص]", encoding="utf-8")
    inner_inventory[fid] = {"type": ext.lstrip("."), "chars": len(text or "")}
    processed += 1

(EXC / "_inner_inventory.json").write_text(
    json.dumps(inner_inventory, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"عولج {processed} ملفاً ثنائياً.")
tot_inner = sum(v.get("inner_count", 0) for v in inner_inventory.values() if v.get("type") == "archive")
print(f"إجمالي الملفات الداخلية المستخرجة من الأرشيفات: {tot_inner}")
print(f"جرد داخلي: {EXC / '_inner_inventory.json'}")
