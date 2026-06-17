#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
auto_index.py — فهرسة تلقائية للملفات غير المفهرسة.

إن لم تتوفّر بيانات وصفية (staging/parts/*.jsonl) ولا فهرس دراسة (parts_clean)،
يبني المحرّك الفهرس تلقائياً من الملفات نفسها:
  - من مجلد المستندات staging/docs/  (ملفات حقيقية: pdf/docx/xlsx/...) فيستخرج نصّها
    إلى staging/text/ ويفهرسها بعناوينها الحقيقية.
  - أو من ملفات النص الجاهزة staging/text/*.txt (يفهرسها باسم الملف).
يكتب: staging/parts/auto_index.jsonl
الاستخدام: python3 auto_index.py <staging_dir> [--force]
"""
import sys
import re
import json
import hashlib
from pathlib import Path

STAGING = Path(sys.argv[1] if len(sys.argv) > 1 else "staging")
FORCE = "--force" in sys.argv
TEXT = STAGING / "text"
PARTS = STAGING / "parts"
DOCS = STAGING / "docs"
PARTS.mkdir(parents=True, exist_ok=True)
TEXT.mkdir(parents=True, exist_ok=True)

DOC_EXT = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".txt"}


def already_indexed():
    if (STAGING / "parts_clean").exists() and any((STAGING / "parts_clean").glob("*.jsonl")):
        return True
    for jl in PARTS.glob("*.jsonl"):
        if jl.stat().st_size > 2:
            return True
    return False


def _id(name):
    return "doc_" + hashlib.md5(name.encode("utf-8")).hexdigest()[:16]


def main():
    if already_indexed() and not FORCE:
        print("الملفات مفهرسة مسبقاً — لا حاجة للفهرسة التلقائية (استخدم --force للإجبار).")
        return

    def read_inner(p):
        """استخراج نص مبسّط (دون استيراد process_binaries لتفادي تنفيذه الذاتي)."""
        ext = p.suffix.lower()
        try:
            if ext in (".txt", ".csv"):
                return p.read_text(encoding="utf-8", errors="replace")
            if ext == ".docx":
                import docx
                return "\n".join(x.text for x in docx.Document(str(p)).paragraphs)
            if ext in (".xlsx", ".xls"):
                import openpyxl
                wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
                out = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c is not None]
                        if cells:
                            out.append(" | ".join(cells))
                return "\n".join(out)
            if ext == ".pdf":
                import pypdf
                return "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(str(p)).pages)
        except Exception as e:  # noqa
            return f"[تعذّر استخراج {p.name}: {e}]"
        return ""

    records = []

    # 1) ملفات حقيقية في staging/docs/ (عناوين حقيقية)
    if DOCS.exists():
        for p in sorted(DOCS.rglob("*")):
            if not p.is_file() or p.suffix.lower() not in DOC_EXT:
                continue
            fid = _id(p.name)
            text = read_inner(p) if read_inner else (
                p.read_text(encoding="utf-8", errors="replace") if p.suffix.lower() == ".txt" else "")
            (TEXT / f"{fid}.txt").write_text(text or "", encoding="utf-8")
            ok = bool(text and len(text.strip()) >= 20)
            records.append({
                "id": fid, "title": p.name, "mimeType": "", "fileExtension": p.suffix.lstrip("."),
                "fileSize": p.stat().st_size, "createdTime": None, "modifiedTime": None,
                "viewUrl": "", "owner": None, "parentId": "",
                "parent_path": str(p.parent.relative_to(DOCS)) if p.parent != DOCS else "",
                "is_folder": False, "text_chars": len(text or ""),
                "text_file": str(TEXT / f"{fid}.txt"),
                "read_status": "ok" if ok else "empty", "read_error": None,
            })

    # 2) ملفات نص جاهزة في staging/text/ لم تُفهرس
    indexed_ids = {r["id"] for r in records}
    for p in sorted(TEXT.glob("*.txt")):
        if p.name.startswith("_") or p.stem in indexed_ids:
            continue
        text = p.read_text(encoding="utf-8", errors="replace")
        ok = len(text.strip()) >= 20
        records.append({
            "id": p.stem, "title": p.stem, "mimeType": "", "fileExtension": "txt",
            "fileSize": p.stat().st_size, "createdTime": None, "modifiedTime": None,
            "viewUrl": "", "owner": None, "parentId": "", "parent_path": "",
            "is_folder": False, "text_chars": len(text),
            "text_file": str(p), "read_status": "ok" if ok else "empty", "read_error": None,
        })

    out = PARTS / "auto_index.jsonl"
    with open(out, "w", encoding="utf-8") as f:
        for r in records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"فُهرس تلقائياً {len(records)} ملفاً ⇐ {out}")


if __name__ == "__main__":
    main()
